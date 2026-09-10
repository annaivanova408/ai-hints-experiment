"""Convert the approved research matrix and anxiety questionnaire to runtime JSON."""
from __future__ import annotations

import argparse
import itertools
import json
import random
from zipfile import ZipFile
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook


CONDITIONS = ("AI", "EXPERT", "CONTROL")


def condition_sequence(count: int, rng: random.Random, previous: str | None) -> list[str]:
    source = [condition for condition in CONDITIONS for _ in range(count // 3)]
    for _ in range(100):
        rng.shuffle(source)
        if previous and source[0] == previous:
            continue
        if any(source[i] == source[i + 1] == source[i + 2] for i in range(len(source) - 2)):
            continue
        return list(source)
    raise RuntimeError("Could not build a constrained condition sequence")


def rotate_conditions(assignments: dict[str, str], offset: int) -> dict[str, str]:
    return {
        segment_id: CONDITIONS[(CONDITIONS.index(condition) + offset) % len(CONDITIONS)]
        for segment_id, condition in assignments.items()
    }


def build_config(matrix_path: Path, anxiety_path: Path) -> dict:
    workbook = load_workbook(matrix_path, data_only=True)
    montage = {
        str(row[2]): float(row[7])
        for row in workbook["Монтажный лист"].iter_rows(min_row=2, values_only=True)
        if row[2]
    }

    sheet = workbook["Сегменты"]
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    segments = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        item = dict(zip(headers, row))
        clip_file = str(item["Клип (файл)"])
        segments.append(
            {
                "id": str(item["ID"]),
                "video_id": Path(str(item["Видео"])).stem,
                "within_video": int(item["№ в видео"]),
                "clip_file": clip_file,
                "duration_sec": montage[clip_file],
                "topic": str(item["Тема сегмента"]),
                "expected_difficulty": str(item["Ожид. сложность"]),
                "stem": str(item["Вопрос в сегменте"]),
                "options": {
                    "A": str(item["Вариант A"]),
                    "B": str(item["Вариант B"]),
                    "C": str(item["Вариант C"]),
                    "D": str(item["Вариант D"]),
                },
                "correct": str(item["Верный"]).strip(),
                "hints": {
                    "ai": str(item["Подсказка ИИ"]),
                    "expert": str(item["Подсказка Эксперта"]),
                    "control": str(item["Подсказка Контроль (плацебо)"]),
                },
                "control_template": int(item["№ шаблона контроля"]),
            }
        )

    final_test = []
    for row in workbook["Итоговый тест"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        final_test.append(
            {
                "id": str(row[0]),
                "segment_id": str(row[1]),
                "video_id": Path(str(row[2])).stem,
                "topic": str(row[3]),
                "expected_difficulty": str(row[4]),
                "stem": str(row[5]),
                "options": {"A": str(row[6]), "B": str(row[7]), "C": str(row[8]), "D": str(row[9])},
                "correct": str(row[10]).strip(),
            }
        )

    videos = []
    video_titles = {"video_1": "Учебное видео 1", "video_2": "Учебное видео 2", "video_4": "Учебное видео 3"}
    for video_id in ("video_1", "video_2", "video_4"):
        clips = []
        for row in workbook["Монтажный лист"].iter_rows(min_row=2, values_only=True):
            if row[1] and Path(str(row[1])).stem == video_id:
                clips.append(
                    {
                        "file": str(row[2]),
                        "duration_sec": float(row[7]),
                        "segment_id": str(row[9]) if str(row[8]).lower() == "да" else None,
                    }
                )
        videos.append({"id": video_id, "title": video_titles[video_id], "clips": clips})

    with ZipFile(anxiety_path) as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    statements = []
    stimulus_table = document_xml.xpath(".//w:tbl", namespaces=namespace)[2]
    for row in stimulus_table.xpath("./w:tr", namespaces=namespace):
        cells = ["".join(cell.xpath(".//w:t/text()", namespaces=namespace)).strip() for cell in row.xpath("./w:tc", namespaces=namespace)]
        if len(cells) >= 2 and cells[0].isdigit() and 1 <= int(cells[0]) <= 32:
            statements.append({"id": int(cells[0]), "text": cells[1]})
    statements.sort(key=lambda item: item["id"])
    if len(statements) != 32:
        raise ValueError(f"Expected 32 anxiety items, found {len(statements)}")

    video_ids = ["video_1", "video_2", "video_4"]
    video_permutations = list(itertools.permutations(video_ids))
    nasa_permutations = list(itertools.permutations(CONDITIONS))
    assignment_slots = []
    # Build slots in triads. Rotating the three condition labels makes every
    # segment appear exactly ten times in every condition across 30 slots.
    for triad in range(10):
        rng = random.Random(20260909 + triad)
        while True:
            base_conditions = {}
            boundaries = {}
            for video_id in video_ids:
                video_segments = [item for item in segments if item["video_id"] == video_id]
                sequence = condition_sequence(len(video_segments), rng, None)
                base_conditions.update({item["id"]: condition for item, condition in zip(video_segments, sequence)})
                boundaries[video_id] = (sequence[0], sequence[-1])
            triad_orders = [video_permutations[(triad * 3 + offset) % 6] for offset in range(3)]
            if all(boundaries[order[i]][1] != boundaries[order[i + 1]][0] for order in triad_orders for i in range(2)):
                break
        for offset in range(3):
            slot_index = triad * 3 + offset
            assignment_slots.append({
                "index": slot_index,
                "video_order": list(video_permutations[slot_index % 6]),
                "nasa_order": list(nasa_permutations[(slot_index * 5) % 6]),
                "segment_conditions": rotate_conditions(base_conditions, offset),
            })

    return {
        "config_version": "2026-09-10-matrix-v2",
        "video_ids": video_ids,
        "assignment_slots": assignment_slots,
        "settings": {
            "background": "#808080",
            "fixation_ms": 2000,
            "hint_min_ms": 5000,
            "hint_max_ms": 20000,
            "satisfaction_timeout_ms": 10000,
            "recovery_ms": 10000,
            "probe_timeout_ms": 30000,
            "final_test_timeout_ms": 60000,
        },
        "videos": videos,
        "segments": segments,
        "final_test": final_test,
        "nasa": [
            {
                "id": "NASALoad",
                "title": "Умственная нагрузка",
                "text": "Насколько умственно напряжённой была работа в эти моменты? Много ли усилий требовалось на то, чтобы думать, решать, запоминать, искать?",
                "left": "Очень низкая",
                "right": "Очень высокая",
            },
            {
                "id": "NASASuccess",
                "title": "Успешность выполнения",
                "text": "Насколько успешно, на Ваш взгляд, Вы справлялись с заданиями в эти моменты?",
                "left": "Полная неудача",
                "right": "Полный успех",
                "reverse_key": True,
            },
            {
                "id": "NASAEfforts",
                "title": "Усилия",
                "text": "Насколько сильно Вам приходилось стараться, чтобы достичь такого уровня выполнения?",
                "left": "Минимальные",
                "right": "Максимальные",
            },
            {
                "id": "NASAFrustration",
                "title": "Уровень фрустрации",
                "text": "Насколько раздражёнными, неуверенными, обескураженными или напряжёнными Вы себя чувствовали в эти моменты?",
                "left": "Совсем нет",
                "right": "Очень сильно",
            },
        ],
        "anxiety": {
            "instruction": "Ниже приведены утверждения о том, как люди относятся к искусственному интеллекту. Оцените, насколько каждое утверждение соответствует Вам, по шкале от 1 до 7, где 1 - «совершенно не согласен(на)», а 7 - «полностью согласен(на)».",
            "items": statements,
            "scales": {
                "PVA": [1, 2, 3, 4],
                "BBA": [5, 6, 7, 8],
                "JRA": [9, 10, 11, 12],
                "LA": [13, 14, 15, 16],
                "ERA": [17, 18, 19, 20],
                "AEA": [21, 22, 23, 24],
                "ACA": [25, 26, 27, 28],
                "LOTA": [29, 30, 31, 32],
            },
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("matrix", type=Path)
    parser.add_argument("anxiety", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    config = build_config(args.matrix, args.anxiety)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Created {args.output}: {len(config['segments'])} segments, {len(config['final_test'])} final items")


if __name__ == "__main__":
    main()
